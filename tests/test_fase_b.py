"""Mejoras de la v1.2.0: dividendos, exportar a Excel, compras en otra moneda y Linux."""
import datetime as dt
import io
import json
import os
import unittest
from unittest import mock

from comun import RAIZ, ConCarpeta, PreciosInventados, cartera, etf, manual, mov, serie_inventada
from app import exportar, importar, motor, plantilla, servidor


class Dividendos(ConCarpeta):
    def calcula(self):
        hoy = dt.date.today()
        hace = lambda dias: (hoy - dt.timedelta(days=dias)).isoformat()
        cfg = cartera([etf("ko", "KO", tipo="accion"), etf("vti", "VTI.DE")], [
            mov("2023-01-02", "ko", "compra", 10, 1000), mov("2023-01-02", "vti", "compra", 10, 1000),
            mov(hace(500), "ko", "dividendo", 0, 20), mov(hace(200), "ko", "dividendo", 0, 25),
            mov(hace(30), "ko", "dividendo", 0, 15), mov(hace(100), "vti", "dividendo", 0, 10)])
        series = {"KO": serie_inventada(100, "2022-12-01", paso=0), "VTI.DE": serie_inventada(100, "2022-12-01", paso=0)}
        with PreciosInventados(series):
            return motor.construir(cfg, self.carpeta, descargar=False)

    def test_resumen(self):
        dv = self.calcula()["dividendos"]
        self.assertEqual(dv["total"], 70)
        self.assertEqual(dv["ultimos12"], 50)
        self.assertAlmostEqual(dv["rentabilidad"], 50 / 2000)            # sobre lo que valen hoy
        ko = dv["porProducto"][0]
        self.assertEqual((ko["id"], ko["total"], ko["ultimos12"], ko["cobros"]), ("ko", 60, 40, 3))
        self.assertAlmostEqual(ko["sobreCoste"], 40 / 1000)
        self.assertEqual(dv["porAno"][-1][0], str(dt.date.today().year))   # años seguidos hasta hoy
        self.assertEqual(sum(v for _, v in dv["porAno"]), 70)

    def test_sin_dividendos(self):
        cfg = cartera([etf()], [mov("2024-01-02", "etf", "compra", 1, 10)])
        with PreciosInventados():
            dv = motor.construir(cfg, self.carpeta, descargar=False)["dividendos"]
        self.assertEqual((dv["total"], dv["porProducto"], dv["porAno"]), (0, [], []))

    def test_web_sin_importes_esconde_los_dividendos(self):
        datos = self.calcula()
        oculto = exportar.sin_importes(datos)["dividendos"]
        k = oculto["total"] / datos["dividendos"]["total"]
        self.assertNotAlmostEqual(k, 1, places=3)
        self.assertAlmostEqual(oculto["ultimos12"], round(datos["dividendos"]["ultimos12"] * k, 2), places=1)
        self.assertEqual([a for a, _ in oculto["porAno"]], [a for a, _ in datos["dividendos"]["porAno"]])
        self.assertEqual(oculto["rentabilidad"], datos["dividendos"]["rentabilidad"])
        self.assertNotIn(str(datos["dividendos"]["porProducto"][0]["total"]), json.dumps(oculto["porProducto"]))


class ExportarExcel(ConCarpeta):
    def test_se_puede_volver_a_importar(self):
        cfg = cartera([etf(identificador="IE0000000001"), manual("cuenta", "efectivo", nombre="Cuenta del banco")],
                      [mov("2024-01-02", "etf", "compra", 2, 40, comision=1, nota="primera"),
                       mov("2024-03-01", "etf", "dividendo", 0, 3)],
                      [{"id": "v1", "fecha": "2024-03-31", "producto": "cuenta", "valor": 2500}])
        xlsx = plantilla.mis_datos(cfg)
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(xlsx))
        self.assertEqual(libro.sheetnames, ["Movimientos", "Productos"])
        self.assertEqual(libro["Productos"].max_row, 3)
        filas, error = importar.leer_tabla("mis_datos.xlsx", xlsx)
        self.assertIsNone(error)
        self.assertEqual(len(filas), 3)
        with PreciosInventados({"TEST.DE": serie_inventada(20, "2024-01-01", paso=0)}):
            informe = importar.vista_previa(cfg, importar.preparar_tabla(cfg, filas, self.carpeta))
        self.assertEqual(informe["errores"], [])
        self.assertEqual(informe["repetidos"], 2)        # ya estaban: no se duplican
        self.assertEqual(informe["productosNuevos"], [])

    def test_endpoint(self):
        with mock.patch.object(servidor, "DATOS", self.carpeta):
            r = servidor.app.test_client().get("/api/exportar-excel")
        self.assertEqual(r.status_code, 200)
        self.assertIn("rumbo_mis_datos_", r.headers["Content-Disposition"])
        self.assertTrue(r.data.startswith(b"PK"))       # un .xlsx es un zip


class OtraMoneda(ConCarpeta):
    def setUp(self):
        super().setUp()
        fx = serie_inventada(0.9, "2024-01-01", paso=0)
        self.parches = [mock.patch.object(servidor, "DATOS", self.carpeta),
                        mock.patch.object(motor, "serie_cambio", lambda moneda, carpeta, fresca=False:
                                          fx if moneda == "USD" else {})]
        for p in self.parches:
            p.start()
        self.web = servidor.app.test_client()
        self.web.post("/api/empezar", json={"desde": "vacia"})
        cfg = json.load(open(os.path.join(self.carpeta, "cartera.json"), encoding="utf-8"))
        cfg["productos"].append(etf("aapl", "AAPL", "USD", tipo="accion"))
        json.dump(cfg, open(os.path.join(self.carpeta, "cartera.json"), "w", encoding="utf-8"))

    def tearDown(self):
        for p in self.parches:
            p.stop()
        super().tearDown()

    def test_compra_en_dolares_se_guarda_en_euros(self):
        with PreciosInventados():
            r = self.web.post("/api/movimientos", json={"producto": "aapl", "tipo": "compra", "fecha": "2025-04-02",
                                                        "unidades": "5", "importe": "1.000", "comision": "2",
                                                        "moneda": "USD", "nota": "Mi compra"})
        j = r.get_json()
        self.assertEqual(r.status_code, 200, j)
        self.assertEqual(j["item"]["importe"], 900)
        self.assertEqual(j["item"]["comision"], 1.8)
        self.assertEqual(j["item"]["nota"], "Mi compra · 1.000 USD (1 USD = 0,9 €)")

    def test_sin_cambio_ese_dia(self):
        r = self.web.post("/api/movimientos", json={"producto": "aapl", "tipo": "compra", "fecha": "2025-04-02",
                                                    "unidades": "5", "importe": "1000", "moneda": "CHF"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("No encuentro el cambio de CHF", r.get_json()["errores"][0])

    def test_en_euros_no_cambia_nada(self):
        with PreciosInventados():
            j = self.web.post("/api/movimientos", json={"producto": "aapl", "tipo": "compra", "fecha": "2025-04-02",
                                                        "unidades": "5", "importe": "1000", "moneda": "EUR"}).get_json()
        self.assertEqual(j["item"]["importe"], 1000)
        self.assertNotIn("nota", j["item"])

    def test_consulta_del_cambio(self):
        self.assertEqual(self.web.get("/api/cambio?moneda=USD&fecha=2025-04-02").get_json(), {"ok": True, "cambio": 0.9})
        self.assertEqual(self.web.get("/api/cambio?moneda=usd&fecha=hoy").status_code, 400)


class VerEnOtraMoneda(ConCarpeta):
    def test_tipos_de_cambio_de_hoy(self):
        cfg = cartera([etf()], [mov("2024-01-02", "etf", "compra", 1, 10)])
        with PreciosInventados({"USDEUR=X": serie_inventada(0.8, "2024-01-01", paso=0)}):
            m = motor.construir(cfg, self.carpeta, descargar=False)["monedas"]
        self.assertEqual(set(m), set(motor.MONEDAS_VISTA))
        self.assertAlmostEqual(m["USD"]["porEuro"], 1.25)          # 1 € = 1 / 0,8 $
        self.assertEqual(m["USD"]["nombre"], "Dólares")
        self.assertLessEqual((dt.date.today() - dt.date.fromisoformat(m["USD"]["fecha"])).days, 3)


class Linux(unittest.TestCase):
    def test_lanzador(self):
        with open(os.path.join(RAIZ, "Iniciar.sh"), "rb") as f:
            sh = f.read()
        self.assertTrue(sh.startswith(b"#!/bin/bash"))
        self.assertNotIn(b"\r", sh)
        self.assertIn(b"uv run --quiet python -m app", sh)


if __name__ == "__main__":
    unittest.main()
